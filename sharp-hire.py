import streamlit as st
import pandas as pd
import json
import os
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication
from anthropic import Anthropic
from openai import OpenAI
from pypdf import PdfReader
from docx import Document
from fpdf import FPDF

# ==============================================================================
# 🧠 SHARP-STANDARDS PROTOCOL (v2.3.1)
# ==============================================================================

APP_VERSION = "v2.3.1"
st.set_page_config(page_title="Sharp Hire", page_icon="🎯", layout="wide")

# --- CSS: SHARP PALETTE ---
st.markdown("""
<style>
    .stApp { background-color: #0e1117; color: #e0e0e0; }
    .stTextArea textarea, .stTextInput input, .stSelectbox div[data-baseweb="select"] {
        background-color: #1c1c1c !important;
        color: #00e5ff !important;
        border: 1px solid #333 !important;
        font-family: 'Helvetica Neue', sans-serif !important;
    }
    div[data-testid="stFileUploader"] section {
        background-color: #161b22;
        border: 2px dashed #00e5ff; 
        border-radius: 10px;
        min-height: 120px !important; 
        display: flex; align-items: center; justify-content: center;
    }
    div[data-testid="stFileUploader"] section:hover { border-color: #00ffab; }
    h1, h2, h3 {
        background: -webkit-linear-gradient(45deg, #00e5ff, #d500f9);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-weight: 700 !important;
    }
    div[data-testid="stButton"] button {
        background: linear-gradient(45deg, #00e5ff, #00ffab) !important;
        color: #000000 !important;
        border: none !important;
        font-weight: 800 !important;
        text-transform: uppercase;
        transition: transform 0.2s;
    }
    div[data-testid="stButton"] button:hover {
        transform: scale(1.02);
        box-shadow: 0 0 15px #00ffab;
    }
    div[data-testid="stMetricValue"] {
        color: #39ff14 !important; 
        font-family: monospace;
        font-size: 1.4rem !important;
    }
    .status-box {
        background-color: #1c1c1c;
        border-left: 3px solid #00e5ff;
        padding: 10px;
        font-family: monospace;
        color: #aaa;
        font-size: 0.9rem;
    }
    .stAlert { background-color: #1c1c1c; border: 1px solid #333; color: #00e5ff; }
</style>
""", unsafe_allow_html=True)

# --- SESSION STATE ---
if 'candidates_list' not in st.session_state: st.session_state.candidates_list = []
if 'jd_text' not in st.session_state: st.session_state.jd_text = ""
if 'processing_log' not in st.session_state: st.session_state.processing_log = "Ready."
if 'total_cost' not in st.session_state: st.session_state.total_cost = 0.0
if 'costs' not in st.session_state: st.session_state.costs = {"OpenAI (Audio)": 0.0, "Anthropic (Intel)": 0.0}

# --- SECRETS ---
try:
    ANTHROPIC_API_KEY = st.secrets.get("ANTHROPIC_API_KEY") or os.environ.get("ANTHROPIC_API_KEY")
    OPENAI_API_KEY = st.secrets.get("OPENAI_API_KEY") or os.environ.get("OPENAI_API_KEY")
except:
    st.error("❌ Missing AI API Keys. Check secrets.toml")
    st.stop()

anthropic_client = Anthropic(api_key=ANTHROPIC_API_KEY)
openai_client = OpenAI(api_key=OPENAI_API_KEY)

# --- UTILITIES ---

def update_status(msg):
    st.session_state.processing_log = msg

def track_cost(provider, amount):
    st.session_state.costs[provider] += amount
    st.session_state.total_cost += amount

def extract_text_from_file(file):
    try:
        file_type = file.name.split('.')[-1].lower()
        if file_type in ['mp3', 'm4a', 'wav', 'mp4', 'mpeg', 'mpga']:
            return transcribe_audio(file)
        elif file_type == 'pdf':
            reader = PdfReader(file)
            return "\n".join([page.extract_text() for page in reader.pages])
        elif file_type == 'docx':
            doc = Document(file)
            return "\n".join([para.text for para in doc.paragraphs])
        elif file_type in ['txt', 'md']:
            return file.read().decode("utf-8")
        return "Unsupported format."
    except Exception as e:
        return f"Error extracting {file.name}: {str(e)}"

def transcribe_audio(file):
    try:
        transcript = openai_client.audio.transcriptions.create(model="whisper-1", file=file)
        track_cost("OpenAI (Audio)", 0.06) 
        return transcript.text
    except Exception as e:
        return f"Whisper Error: {str(e)}"

def clean_json_response(txt):
    txt = txt.strip()
    if "```json" in txt: txt = txt.split("```json")[1].split("```")[0]
    elif "```" in txt: txt = txt.split("```")[1].split("```")[0]
    return txt.strip()

# --- PDF GENERATOR ---
class SharpPDF(FPDF):
    def header(self):
        self.set_fill_color(14, 17, 23) # Deep Black
        self.rect(0, 0, 210, 40, 'F')
        self.set_font('Arial', 'B', 16)
        self.set_text_color(0, 229, 255) # Neon Cyan
        self.cell(0, 10, 'SHARP HIRE | INTELLIGENCE REPORT', 0, 1, 'C')
        self.ln(10)

    def section_title(self, label):
        self.set_font('Arial', 'B', 12)
        self.set_text_color(0, 229, 255) # Cyan
        self.cell(0, 10, label, 0, 1, 'L')
        self.set_text_color(0, 0, 0) # Black for body

    def chapter_body(self, body):
        self.set_font('Arial', '', 10)
        self.multi_cell(0, 6, body)
        self.ln()

def generate_sharp_pdf(results):
    pdf = SharpPDF()
    pdf.add_page()
    pdf.section_title("SESSION SUMMARY")
    pdf.chapter_body(f"Candidates Analyzed: {len(results)}")
    pdf.ln(5)
    
    for res in results:
        cand = res['candidate']
        rec = res['recruiter']
        pdf.set_fill_color(240, 240, 240)
        pdf.rect(10, pdf.get_y(), 190, 10, 'F')
        pdf.section_title(f"CANDIDATE: {cand['name']}  (Verdict: {cand['verdict']})")
        
        # Safe score access
        s = cand['scores']
        scores = f"Paper Fit: {s.get('cv_match_score',0)}/10 | Actual Fit: {s.get('interview_performance_score',0)}/10 | Truth: {s.get('cv_truthfulness',0)}/10"
        pdf.chapter_body(scores)
        
        pdf.chapter_body(f"Summary: {res['executive_summary']}")
        pdf.ln(5)
        
    return pdf.output(dest='S').encode('latin-1', 'replace')

# --- EMAIL ENGINE ---
def send_email(to_email, pdf_bytes):
    sender_email = st.secrets.get("EMAIL_USER")
    sender_password = st.secrets.get("EMAIL_PASSWORD")
    smtp_host = st.secrets.get("SMTP_HOST", "smtp.gmail.com")
    smtp_port = st.secrets.get("SMTP_PORT", 587)

    if not (sender_email and sender_password):
        return "❌ SMTP Secrets Missing (EMAIL_USER, EMAIL_PASSWORD)."

    try:
        msg = MIMEMultipart()
        msg['From'] = sender_email
        msg['To'] = to_email
        msg['Subject'] = "Sharp Hire Intelligence Report"
        body = "Attached is the forensic interview analysis report from Sharp Hire."
        msg.attach(MIMEText(body, 'plain'))
        attachment = MIMEApplication(pdf_bytes, Name="Sharp_Hire_Report.pdf")
        attachment['Content-Disposition'] = 'attachment; filename="Sharp_Hire_Report.pdf"'
        msg.attach(attachment)
        with smtplib.SMTP(smtp_host, smtp_port) as server:
            server.starttls()
            server.login(sender_email, sender_password)
            server.send_message(msg)
        return "✅ Email Sent Successfully!"
    except Exception as e:
        return f"Error sending email: {str(e)}"

# --- ANALYSIS ENGINE ---
def analyze_comprehensive(transcript, cv_text, jd_text):
    # FIX: Explicitly demanding 'cv_truthfulness' in the prompt
    system_prompt = f"""
    You are a FORENSIC Talent Auditor. Perform a deep, multi-vector analysis.
    
    **DATA:** JD (Required), CV (Claims), TRANSCRIPT (Evidence).

    **ANALYSIS VECTORS:**
    1. **Recruiter:** Did they dig deep?
    2. **Cand vs JD:** Skills match?
    3. **Cand vs Questions:** Answer Quality/Directness.
    4. **Cand vs CV:** Truthfulness (Did they lie?).

    **OUTPUT JSON STRUCTURE:**
    {{
        "executive_summary": "High-level narrative.",
        "candidate": {{
            "name": "Inferred Name",
            "scores": {{ 
                "cv_match_score": 0, 
                "interview_performance_score": 0, 
                "technical_depth": 0, 
                "culture_fit": 0,
                "cv_truthfulness": 0 
            }},
            "fit_analysis": {{ "gap_analysis": "...", "jd_vs_transcript": "..." }},
            "strengths": ["..."],
            "red_flags": ["..."],
            "verdict": "Hire / No Hire"
        }},
        "recruiter": {{
            "scores": {{ "question_quality": 0, "jd_coverage": 0 }},
            "missed_opportunities": ["..."],
            "coaching_tip": "..."
        }}
    }}
    """
    user_msg = f"JD: {jd_text[:10000]}\nCV: {cv_text[:10000]}\nTRANSCRIPT: {transcript[:40000]}"
    try:
        message = anthropic_client.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=4000,
            temperature=0.1, 
            system=system_prompt,
            messages=[{"role": "user", "content": user_msg}]
        )
        track_cost("Anthropic (Intel)", 0.03) 
        return json.loads(clean_json_response(message.content[0].text))
    except Exception as e:
        return {"error": str(e)}

def render_neon_progress(label, score, max_score=10):
    # Handle missing keys gracefully just in case
    if score is None: score = 0
    
    pct = (score / max_score) * 100
    color = "#ff4b4b"
    if score >= 5: color = "#ffa700"
    if score >= 7: color = "#39ff14"
    if score >= 9: color = "#00e5ff"
    st.markdown(f"""
    <div style="margin-bottom: 8px;">
        <div style="display: flex; justify-content: space-between; font-size: 0.85rem;">
            <span style="color: #bbb;">{label}</span>
            <span style="color: {color}; font-weight: bold;">{score}/10</span>
        </div>
        <div style="background-color: #222; height: 6px; border-radius: 4px;">
            <div style="background-color: {color}; width: {pct}%; height: 100%; border-radius: 4px;"></div>
        </div>
    </div>
    """, unsafe_allow_html=True)

# --- LAYOUT ---

c_title, c_meta = st.columns([3, 1])
with c_title:
    st.title("🎯 Sharp Hire")
    st.caption("Multi-Candidate Interview Intelligence")
with c_meta:
    st.markdown(f"<div style='text-align: right; color: #666;'>{APP_VERSION}</div>", unsafe_allow_html=True)
    st.metric("Session Cost", f"${st.session_state.total_cost:.4f}")
    st.markdown(f"<div class='status-box'><span style='color: #00e5ff;'>● SYSTEM ACTIVE</span><br>{st.session_state.processing_log}</div>", unsafe_allow_html=True)

c1, c2, c3 = st.columns(3)
with c1:
    st.markdown("### 1. The Job")
    jd_file = st.file_uploader("JD (Stays for session)", type=['pdf','docx','txt'], key="jd", label_visibility="collapsed")
with c2:
    st.markdown("### 2. The Candidate")
    cv_file = st.file_uploader("CV (Updates per run)", type=['pdf','docx','txt'], key="cv", label_visibility="collapsed")
with c3:
    st.markdown("### 3. The Interview")
    call_file = st.file_uploader("Audio/Transcript", type=['mp3','wav','m4a','pdf','docx','txt'], key="call", label_visibility="collapsed")

c_btn, c_clear = st.columns([3, 1])
with c_btn:
    start_btn = st.button("Start Forensic Audit (Add to Session)", type="primary", use_container_width=True)
with c_clear:
    if st.button("Reset Session"):
        st.session_state.candidates_list = []
        st.rerun()

# --- PROCESSING ---
if start_btn:
    if not (jd_file and cv_file and call_file):
        st.warning("⚠️ Upload JD, CV, and Transcript.")
    else:
        try:
            with st.status("🚀 Analyzing Candidate...", expanded=True) as status:
                if not st.session_state.jd_text:
                    update_status("Reading JD...")
                    st.session_state.jd_text = extract_text_from_file(jd_file)
                
                update_status("Reading CV & Transcript...")
                cv_txt = extract_text_from_file(cv_file)
                trans_txt = extract_text_from_file(call_file)
                
                update_status("Running Forensic Logic...")
                res = analyze_comprehensive(trans_txt, cv_txt, st.session_state.jd_text)
                
                if "error" not in res:
                    st.session_state.candidates_list.append(res)
                    status.update(label="✅ Added to Session!", state="complete", expanded=False)
                else:
                    st.error(res['error'])
        except Exception as e:
            st.error(f"Error: {e}")

# --- DASHBOARD ---
if st.session_state.candidates_list:
    st.divider()
    st.subheader(f"📊 Assessment Session ({len(st.session_state.candidates_list)} Candidates)")
    
    tabs = st.tabs([f"👤 {c['candidate']['name']}" for c in st.session_state.candidates_list])
    
    for i, tab in enumerate(tabs):
        data = st.session_state.candidates_list[i]
        cand = data['candidate']
        rec = data['recruiter']
        
        with tab:
            st.info(f"**Executive Summary:** {data['executive_summary']}")
            
            col_a, col_b = st.columns(2)
            with col_a:
                st.markdown("#### Candidate Performance")
                with st.container(border=True):
                    s = cand['scores']
                    # Corrected keys match the prompt now
                    render_neon_progress("Paper Fit (CV)", s.get('cv_match_score', 0))
                    render_neon_progress("Actual Fit (Interview)", s.get('interview_performance_score', 0))
                    render_neon_progress("Tech Depth", s.get('technical_depth', 0))
                    render_neon_progress("Truthfulness", s.get('cv_truthfulness', 0))
                
                with st.expander("Details & Flags"):
                    st.write(cand['fit_analysis']['gap_analysis'])
                    for f in cand['red_flags']: st.warning(f)

            with col_b:
                st.markdown("#### Recruiter Performance")
                with st.container(border=True):
                    rs = rec['scores']
                    render_neon_progress("Question Quality", rs.get('question_quality', 0))
                    render_neon_progress("JD Coverage", rs.get('jd_coverage', 0))
                    st.caption(f"Coach: {rec['coaching_tip']}")
                    
    st.divider()
    st.markdown("### 📤 Export Session")
    
    pdf_bytes = generate_sharp_pdf(st.session_state.candidates_list)
    c_down, c_email = st.columns(2)
    
    with c_down:
        st.download_button(
            label="⬇️ Download Full PDF Report",
            data=pdf_bytes,
            file_name="Sharp_Hire_Report.pdf",
            mime="application/pdf",
            use_container_width=True
        )
        
    with c_email:
        email_target = st.text_input("Email Report To:", placeholder="recruiter@company.com")
        if st.button("📧 Send Email"):
            if email_target:
                with st.spinner("Sending..."):
                    res = send_email(email_target, pdf_bytes)
                    if "Success" in res: st.success(res)
                    else: st.error(res)
